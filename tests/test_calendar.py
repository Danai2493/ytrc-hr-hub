"""Company calendar workbook import (Admin > Holidays)."""
import io
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from sqlalchemy import select

from app.models import Calendar, Holiday
from app.services.calendar_import import parse_workbook

YELLOW = PatternFill("solid", fgColor="FFFFFF00")
ORANGE = PatternFill("solid", fgColor="FFFFC000")
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]


def make_workbook(year: int, office_hol: set, factory_hol: set) -> bytes:
    """Same layout as HR's file: Office grid from column B, Factory from column Z,
    yellow = holiday, orange = weekly rest (Office Sunday, Factory Friday)."""
    wb = Workbook()
    wb.active.title = "Manpower"
    ws = wb.create_sheet(f"Calendar {year % 100}")
    ws.cell(1, 2, f"Calendar for year {year} (YTRC) ")
    for grid_col, holidays, rest in ((2, office_hol, 6), (26, factory_hol, 4)):
        for m in range(1, 13):
            top = 2 + (m - 1) // 3 * 9
            col = grid_col + (m - 1) % 3 * 8
            ws.cell(top, col, MONTHS[m - 1])
            for k, dn in enumerate(["Mo", "Tu", "We", "Th", "Fr", "Sa", "Su"]):
                ws.cell(top + 1, col + k, dn)
            d = date(year, m, 1)
            week = 0
            while d.month == m:
                c = ws.cell(top + 2 + week, col + d.weekday(), d.day)
                if d in holidays:
                    c.fill = YELLOW
                elif d.weekday() == rest:
                    c.fill = ORANGE
                if d.weekday() == 6:
                    week += 1
                d += timedelta(days=1)
    # the list beside the grids; the year typed there is stale on purpose
    ws.cell(2, 51, "Holidays and Observances:")
    for i, d in enumerate(sorted(office_hol | factory_hol)):
        ws.cell(4 + i, 52, d.replace(year=year - 2))
        ws.cell(4 + i, 53, f"Holiday {d:%d %b}")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_reads_colours_not_list_years():
    office = {date(2028, 1, 1), date(2028, 4, 13)}
    factory = {date(2028, 4, 13), date(2028, 10, 22)}
    p = parse_workbook(make_workbook(2028, office, factory))
    assert p.year == 2028
    assert p.rest == {"OFFICE": [6], "FACTORY": [4]}
    assert [d for d, _ in p.holidays["OFFICE"]] == sorted(office)
    assert [d for d, _ in p.holidays["FACTORY"]] == sorted(factory)
    assert p.holidays["FACTORY"][1][1] == "Holiday 22 Oct"


def test_admin_upload_replaces_year(login, db):
    hr = login("hr.manager")
    content = make_workbook(2029, {date(2029, 1, 1)}, {date(2029, 1, 2)})
    r = hr.post("/admin/calendar/preview", files={"file": ("Calendar 2029.xlsx", content, "application/octet-stream")})
    assert r.status_code == 200 and "Holiday 01 Jan" in r.text
    import html
    import re

    data = html.unescape(re.search(r'name="data" value="([^"]+)"', r.text).group(1))
    hr.post("/admin/calendar/apply", {"data": data})
    db.expire_all()
    days = db.scalars(select(Holiday.day).where(Holiday.calendar_code == "FACTORY", Holiday.day >= date(2029, 1, 1))).all()
    assert days == [date(2029, 1, 2)]
    assert db.get(Calendar, "FACTORY").rest_weekdays == "4"
