"""Read the company calendar workbook HR keeps each year ("Calendar for year
2026 (YTRC)") and load it as the Office and Factory holiday calendars.

The workbook has two month grids side by side: Office on the left, Factory on
the right. Day colours carry the meaning: yellow = holiday, orange = weekly
rest day. Holiday names come from the lists next to the grids, matched by
month and day (the years typed in those lists are often left over from an
older file, so only month and day are trusted)."""
from __future__ import annotations

import io
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime

from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import Calendar, Holiday
from .common import audit

MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]
HOLIDAY_FILL = "FFFF00"
REST_FILL = "FFC000"
FACTORY_FROM_COL = 26  # column Z: the Factory grid starts here
LIST_FROM_COL = 51  # column AY: "Holidays and Observances" list
FLAG_COLS = {"OFFICE": 58, "FACTORY": 59}  # BF / BG: 1 = holiday on that calendar


class CalendarFileError(ValueError):
    pass


@dataclass
class ParsedCalendar:
    year: int
    rest: dict[str, list[int]] = field(default_factory=dict)  # calendar -> weekday numbers, Monday = 0
    holidays: dict[str, list[tuple[date, str]]] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)

    def to_json(self) -> dict:
        return {
            "year": self.year,
            "rest": self.rest,
            "holidays": {k: [[d.isoformat(), n] for d, n in v] for k, v in self.holidays.items()},
            "warnings": self.warnings,
        }

    @classmethod
    def from_json(cls, data: dict) -> "ParsedCalendar":
        return cls(
            year=int(data["year"]),
            rest={k: [int(x) for x in v] for k, v in data["rest"].items()},
            holidays={k: [(date.fromisoformat(d), n) for d, n in v] for k, v in data["holidays"].items()},
            warnings=list(data.get("warnings") or []),
        )


def _fill(cell) -> str | None:
    f = cell.fill
    if not f or not f.fill_type:
        return None
    rgb = f.fgColor.rgb
    return rgb[-6:].upper() if isinstance(rgb, str) else None


def parse_workbook(content: bytes) -> ParsedCalendar:
    try:
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception as exc:  # not an xlsx, or damaged
        raise CalendarFileError("This is not a readable .xlsx file") from exc
    sheets = [ws for ws in wb.worksheets if ws.title.lower().startswith("calendar")]
    if not sheets:
        raise CalendarFileError("No sheet named 'Calendar ..' in the file")
    ws = sheets[0]
    year = None
    for row in ws.iter_rows(min_row=1, max_row=3):
        for c in row:
            if isinstance(c.value, str) and "year" in c.value.lower():
                digits = "".join(ch for ch in c.value if ch.isdigit())
                if len(digits) == 4:
                    year = int(digits)
                    break
        if year:
            break
    if year is None:
        tail = "".join(ch for ch in ws.title if ch.isdigit())
        year = 2000 + int(tail) if len(tail) == 2 else None
    if year is None:
        raise CalendarFileError("Cannot find the year in the sheet title")

    days: dict[tuple[str, date], str | None] = {}
    names: dict[str, dict[tuple[int, int], str]] = {"OFFICE": {}, "FACTORY": {}, "LIST": {}}
    flagged: dict[str, set[tuple[int, int]]] = {"OFFICE": set(), "FACTORY": set()}
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for c in row:
            v = c.value
            # a month grid: month name, then a row of day names starting "Mo"
            if isinstance(v, str) and v.strip() in MONTHS and ws.cell(c.row + 1, c.column).value == "Mo":
                m = MONTHS.index(v.strip()) + 1
                grid = "OFFICE" if c.column < FACTORY_FROM_COL else "FACTORY"
                for r in range(c.row + 2, c.row + 8):
                    for k in range(7):
                        cell = ws.cell(r, c.column + k)
                        n = cell.value
                        if isinstance(n, (int, float)) and 1 <= n <= 31:
                            try:
                                d = date(year, m, int(n))
                            except ValueError:
                                continue
                            if d.weekday() == k:
                                days[(grid, d)] = _fill(cell)
            # a holiday list line: a date, then its name a few cells to the right
            elif isinstance(v, (datetime, date)):
                name = None
                for k in range(1, 5):
                    t = ws.cell(c.row, c.column + k).value
                    if isinstance(t, str) and t.strip():
                        name = " ".join(t.split())
                        break
                if not name:
                    continue
                key = (v.month, v.day)
                where = "LIST" if c.column >= LIST_FROM_COL else ("OFFICE" if c.column < FACTORY_FROM_COL else "FACTORY")
                names[where].setdefault(key, name)
                if where == "LIST":
                    for cal, col in FLAG_COLS.items():
                        if ws.cell(c.row, col).value in (1, "1"):
                            flagged[cal].add(key)
    if not days:
        raise CalendarFileError("No month grids found in the sheet")

    out = ParsedCalendar(year=year)
    for cal in ("OFFICE", "FACTORY"):
        cal_days = {d: f for (g, d), f in days.items() if g == cal}
        if len(cal_days) < 365:
            out.warnings.append(f"{cal}: only {len(cal_days)} days found in the grid")
        rest = Counter(d.weekday() for d, f in cal_days.items() if f == REST_FILL)
        # a weekday is a rest day when nearly every week shows it in orange
        out.rest[cal] = sorted(w for w, n in rest.items() if n >= 40)
        hol = []
        for d in sorted(d for d, f in cal_days.items() if f == HOLIDAY_FILL):
            key = (d.month, d.day)
            name = names["LIST"].get(key) or names[cal].get(key)
            if not name:
                name = "Company holiday"
                out.warnings.append(f"{cal} {d:%d/%m/%Y}: yellow in the grid but no name in the lists")
            hol.append((d, name[:120]))
        out.holidays[cal] = hol
        marked = {(d.month, d.day) for d, _ in hol}
        for m, dd in sorted(flagged[cal] - marked):
            if date(year, m, dd).weekday() in out.rest[cal]:
                continue  # falls on the weekly rest day, so nothing to add
            out.warnings.append(f"{cal} {dd:02d}/{m:02d}: ticked in the list but not yellow in the grid (grid used)")
    return out


def apply(db: Session, parsed: ParsedCalendar, user) -> dict[str, int]:
    """Replace one year's holidays on both calendars and set the weekly rest days."""
    y0, y1 = date(parsed.year, 1, 1), date(parsed.year, 12, 31)
    counts = {}
    for cal_code, hol in parsed.holidays.items():
        cal = db.get(Calendar, cal_code)
        if cal is None:
            continue
        rest = ",".join(str(w) for w in parsed.rest.get(cal_code, []))
        if rest and rest != cal.rest_weekdays:
            audit(db, user, "update", "calendar", cal_code, f"rest days {cal.rest_weekdays} -> {rest}")
            cal.rest_weekdays = rest
        old = db.scalars(select(Holiday).where(Holiday.calendar_code == cal_code, Holiday.day >= y0, Holiday.day <= y1)).all()
        for h in old:
            db.delete(h)
        db.flush()
        for d, name in hol:
            db.add(Holiday(calendar_code=cal_code, day=d, name=name))
        counts[cal_code] = len(hol)
        audit(db, user, "import", "holiday", f"{cal_code}:{parsed.year}", f"{len(old)} replaced by {len(hol)}")
    db.flush()
    return counts
